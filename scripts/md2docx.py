#!/usr/bin/env python3
"""
Markdown → Word (.docx) 转换器（面向含大量表格与 ASCII 线框的技术文档）

用法：
    .venv/bin/python scripts/md2docx.py <输入.md> [输出.docx]

为什么自己解析而不用 textutil 转 HTML：
    实测 `textutil -convert docx` **不生成真正的 Word 表格**（`<w:tbl>` 数量为 0），
    它会把表格拍平成纯文本行。对一份靠 49 个表格组织信息的 PRD 来说，
    这等于把结构毁掉。所以用 python-docx 直接构造，表格、线框都能精确控制。

三条硬要求（都踩过坑）：
    1. 表格必须是**真表格**，保留表头与边框
    2. ASCII 线框必须用**等宽字体**且**逐行保留空格**，不能被重排
    3. 中文正文用中文字体，代码/线框用等宽字体，两者不能混
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

CJK_FONT = 'PingFang SC'
MONO_FONT = 'Menlo'
FALLBACK_MONO = 'Courier New'

HEADING = re.compile(r'^(#{1,6})\s+(.+?)\s*$')
BULLET = re.compile(r'^(\s*)[-*+]\s+(.+)$')
NUMBERED = re.compile(r'^(\s*)\d+[.、]\s+(.+)$')
TABLE_ROW = re.compile(r'^\|(.+)\|\s*$')
TABLE_SEP = re.compile(r'^\|[\s:|-]+\|\s*$')
FENCE = re.compile(r'^```')
HR = re.compile(r'^-{3,}\s*$')
BLOCKQUOTE = re.compile(r'^>\s?(.*)$')

INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`)')


def set_cjk_font(run, name: str = CJK_FONT) -> None:
    """让中文也走指定字体（python-docx 默认只管西文）。"""
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rfonts.set(qn(attr), name)


def add_runs(paragraph, text: str, base_size: float | None = None,
             mono: bool = False) -> None:
    """写入一段文本，处理 **粗体** 与 `行内代码`。"""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            set_cjk_font(run, MONO_FONT if mono else CJK_FONT)
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_cjk_font(run, MONO_FONT)
            run.font.size = Pt((base_size or 10.5) - 1)
        else:
            run = paragraph.add_run(part)
            set_cjk_font(run, MONO_FONT if mono else CJK_FONT)
        if base_size:
            run.font.size = Pt(base_size)


def shade(cell, color: str = 'F2F2F2') -> None:
    """给单元格加底色（用于表头）。"""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), color)
    tc_pr.append(shd)


def add_code_block(doc: Document, lines: list[str]) -> None:
    """ASCII 线框 / 代码块：等宽、小字号、逐行保留空格。"""
    for line in lines:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        pf.left_indent = Pt(6)
        # 关键：不加任何自动换行处理，空格原样保留
        run = p.add_run(line if line else ' ')
        set_cjk_font(run, MONO_FONT)
        run.font.size = Pt(7.5)
        run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    """真正的 Word 表格，带表头底色。"""
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j in range(cols):
            text = row[j] if j < len(row) else ''
            cell = cells[j]
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            add_runs(p, text, base_size=8.5)
            if i == 0:
                for run in p.runs:
                    run.bold = True
                shade(cell)
    doc.add_paragraph()


def convert(src: Path, dst: Path) -> int:
    if not src.is_file():
        print(f'❌ 找不到输入文件：{src}')
        return 1

    lines = src.read_text(encoding='utf-8').splitlines()
    doc = Document()

    # 全局默认字体
    style = doc.styles['Normal']
    style.font.name = CJK_FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), CJK_FONT)

    i = 0
    stats = {'tables': 0, 'code_blocks': 0, 'headings': 0}
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # ── 代码块 / 线框 ──
        if FENCE.match(stripped):
            i += 1
            block = []
            while i < len(lines) and not FENCE.match(lines[i].strip()):
                block.append(lines[i])
                i += 1
            i += 1  # 跳过结束的 ```
            add_code_block(doc, block)
            stats['code_blocks'] += 1
            continue

        # ── 表格 ──
        if TABLE_ROW.match(stripped) and i + 1 < len(lines) and TABLE_SEP.match(lines[i + 1].strip()):
            rows = []
            header = [c.strip() for c in TABLE_ROW.match(stripped).group(1).split('|')]
            rows.append(header)
            i += 2
            while i < len(lines):
                s = lines[i].strip()
                m = TABLE_ROW.match(s)
                if not m or TABLE_SEP.match(s):
                    break
                rows.append([c.strip() for c in m.group(1).split('|')])
                i += 1
            add_table(doc, rows)
            stats['tables'] += 1
            continue

        # ── 标题 ──
        m = HEADING.match(stripped)
        if m:
            level = min(len(m.group(1)), 4)
            text = m.group(2)
            h = doc.add_heading('', level=level)
            add_runs(h, text)
            for run in h.runs:
                set_cjk_font(run, CJK_FONT)
                run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)
            stats['headings'] += 1
            i += 1
            continue

        # ── 分隔线（跳过，用段间距代替）──
        if HR.match(stripped):
            i += 1
            continue

        # ── 引用 ──
        if BLOCKQUOTE.match(stripped):
            text = BLOCKQUOTE.match(stripped).group(1)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(14)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, text)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            i += 1
            continue

        # ── 列表 ──
        m = BULLET.match(line)
        if m:
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            add_runs(p, m.group(2))
            i += 1
            continue
        m = NUMBERED.match(line)
        if m:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            add_runs(p, m.group(2))
            i += 1
            continue

        # ── 空行 ──
        if stripped == '':
            i += 1
            continue

        # ── 正文 ──
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        add_runs(p, stripped)
        i += 1

    doc.save(str(dst))
    print(f'  ✅ {dst}  ({dst.stat().st_size / 1024:.0f} KB)')
    print(f'     标题 {stats["headings"]} · 表格 {stats["tables"]} · 线框/代码块 {stats["code_blocks"]}')
    return 0


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix('.docx')
    return convert(src, dst)


if __name__ == '__main__':
    sys.exit(main())
